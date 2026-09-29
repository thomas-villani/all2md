"""Nested and adjacent lists survive a round trip through DOCX.

The renderer used to give every list item the one "List Bullet" or "List Number" style,
whatever its depth, so a nested list flattened into its parent; and every numbered list
shared that style's numbering instance, so Word -- and the parser, which counts as Word
does -- carried the count on from one list into the next and read them as one list.
PDF -> DOCX kept 229 of the 253 lists in the PMC development corpus
(``python -m benchmarks.pmc docx``); the Markdown route kept all of them.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest

from all2md.ast import Document, List, ListItem, Paragraph, Text
from all2md.ast.nodes import Node
from all2md.options.docx import DocxRendererOptions
from all2md.parsers.docx import DocxToAstConverter
from all2md.renderers.docx import DocxRenderer

docx = pytest.importorskip("docx")

pytestmark = [pytest.mark.unit, pytest.mark.docx]


def _item(text: str, *nested: List) -> ListItem:
    return ListItem(children=[Paragraph(content=[Text(content=text)]), *nested])


def _list(ordered: bool, *items: ListItem, start: int = 1) -> List:
    return List(ordered=ordered, items=list(items), start=start)


def _render(document: Document, options: DocxRendererOptions | None = None) -> bytes:
    buffer = BytesIO()
    DocxRenderer(options).render(document, buffer)
    return buffer.getvalue()


def _shape(nodes: list[Node]) -> list[object]:
    """Lists as (ordered, start, [item texts and nested shapes]); anything else by type name."""
    shape: list[object] = []
    for node in nodes:
        if isinstance(node, List):
            items: list[object] = []
            for item in node.items:
                for child in item.children:
                    if isinstance(child, Paragraph):
                        items.append("".join(getattr(inline, "content", "") for inline in child.content))
                    else:
                        items.extend(_shape([child]))
            shape.append((node.ordered, node.start, items))
        else:
            shape.append(type(node).__name__)
    return shape


def _round_trip(document: Document, options: DocxRendererOptions | None = None) -> list[object]:
    return _shape(DocxToAstConverter().parse(BytesIO(_render(document, options))).children)


def _template_without(tmp_path: Path, *styles: str) -> DocxRendererOptions:
    template = docx.Document()
    for name in styles:
        template.styles[name].delete()
    template.save(tmp_path / "template.docx")
    return DocxRendererOptions(template_path=str(tmp_path / "template.docx"))


class TestRenderer:
    def test_a_nested_list_takes_the_style_for_its_depth(self) -> None:
        document = Document(children=[_list(True, _item("a", _list(False, _item("b", _list(True, _item("c"))))))])
        word = docx.Document(BytesIO(_render(document)))
        assert [(p.style.name, p.text) for p in word.paragraphs] == [
            ("List Number", "a"),
            ("List Bullet 2", "b"),
            ("List Number 3", "c"),
        ]

    def test_each_numbered_list_restarts_at_its_start(self) -> None:
        from docx.oxml.ns import qn

        word = docx.Document(
            BytesIO(_render(Document(children=[_list(True, _item("a")), _list(True, _item("b"), start=4)])))
        )
        num_ids = [p._p.pPr.numPr.numId.val for p in word.paragraphs]
        assert num_ids[0] != num_ids[1]
        numbering = word.part.numbering_part.element
        starts = []
        for num_id in num_ids:
            num = next(num for num in numbering.findall(qn("w:num")) if num.get(qn("w:numId")) == str(num_id))
            starts.append(num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")))
        assert starts == ["1", "4"]

    def test_a_depth_the_template_lacks_is_created(self) -> None:
        deep = _list(False, _item("4"))
        for depth in "321":
            deep = _list(False, _item(depth, deep))
        word = docx.Document(BytesIO(_render(Document(children=[deep]))))
        assert word.paragraphs[-1].style.name == "List Bullet 4"
        assert word.paragraphs[-1].style.element.pPr.numPr.numId is not None


class TestRoundTrip:
    def test_adjacent_numbered_lists_stay_apart(self) -> None:
        document = Document(children=[_list(True, _item("a"), _item("b")), _list(True, _item("c"))])
        assert _round_trip(document) == [(True, 1, ["a", "b"]), (True, 1, ["c"])]

    def test_a_nested_list_keeps_its_start_and_its_parents_count(self) -> None:
        document = Document(
            children=[_list(True, _item("a"), _item("b", _list(True, _item("c"), start=5)), _item("d"), start=7)]
        )
        assert _round_trip(document) == [(True, 7, ["a", "b", (True, 5, ["c"]), "d"])]

    @pytest.mark.parametrize("outer", [True, False], ids=["numbered", "bulleted"])
    @pytest.mark.parametrize("inner", [True, False], ids=["numbered", "bulleted"])
    def test_a_nested_list_stays_nested(self, outer: bool, inner: bool) -> None:
        document = Document(children=[_list(outer, _item("a", _list(inner, _item("b"), _item("c"))), _item("d"))])
        assert _round_trip(document) == [(outer, 1, ["a", (inner, 1, ["b", "c"]), "d"])]

    def test_four_levels_deep(self) -> None:
        deep = _list(True, _item("4"))
        for depth in "321":
            deep = _list(True, _item(depth, deep))
        assert _round_trip(Document(children=[deep])) == [
            (True, 1, ["1", (True, 1, ["2", (True, 1, ["3", (True, 1, ["4"])])])])
        ]

    def test_a_template_without_the_list_styles(self, tmp_path: Path) -> None:
        options = _template_without(tmp_path, "List Number", "List Number 2", "List Bullet")
        document = Document(children=[_list(True, _item("a", _list(True, _item("b")))), _list(True, _item("c"))])
        assert _round_trip(document, options) == [(True, 1, ["a", (True, 1, ["b"])]), (True, 1, ["c"])]
