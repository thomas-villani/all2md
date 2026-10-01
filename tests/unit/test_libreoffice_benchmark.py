"""Tests for the LibreOffice corpus lane (`benchmarks.libreoffice`).

The lane is a manual instrument, so these hold its arithmetic and its storage rather than
any reading of the corpus: the word multiset, the scoring of each kind of file, the
paired comparison, the cache naming, and the committed Word reading itself.
"""

from __future__ import annotations

import io
import json
from collections import Counter

import docx
import pytest

from benchmarks.libreoffice import corpus as corpus_module
from benchmarks.libreoffice import sweep as sweep_module
from benchmarks.libreoffice.compare import format_changes, format_report, paired, report_json, score, words
from benchmarks.libreoffice.corpus import (
    SETS,
    CorpusError,
    cache_name,
    core_root,
    corpus_files,
    key_of,
    load_word_reading,
    save_word_reading,
)

pytestmark = [pytest.mark.unit, pytest.mark.docx]


def word_record(main: str, **stories) -> dict:
    return {"status": "ok", "main": main, "footnotes": [], "endnotes": [], "shapes": [], **stories}


def mine(text: str, **extra) -> dict:
    return {"status": "ok", "text": text, "math": [], **extra}


def test_words_drop_word_placeholders_and_ignore_case():
    # A field's delimiters (\x13 \x14 \x15) can sit inside a word; Word's text keeps them.
    assert words("Hello\x13 PAGE \x14World\x15 hello") == Counter({"hello": 2, "page": 1, "world": 1})


@pytest.mark.parametrize("key", ["export/tdf119143.docx", "import/__RefNumPara__.docx"])
def test_cache_names_invert_even_when_the_name_holds_the_separator(key):
    assert key_of(cache_name(key)) == key


def test_scoring_sorts_each_kind_of_file():
    reading = {
        "meta": {},
        "files": {
            "export/same.docx": word_record("Same words here"),
            "export/lost.docx": word_record("Kept lost", footnotes=["Note text"], shapes=["Box text"]),
            "export/fails.docx": word_record("Word opens it"),
            "export/both-fail.docx": {"status": "fail", "error": "encrypted"},
            "export/word-fails.docx": {"status": "fail", "error": "bad"},
            "export/changed.docx": {**word_record("Old"), "sha256": "aaa"},
            "import/unswept.docx": word_record("Never read"),
        },
    }
    sweep = {
        "export/same.docx": mine("same WORDS here"),
        "export/lost.docx": mine("Kept note text box text 1."),
        "export/fails.docx": {"status": "fail", "error": "KeyError: x"},
        "export/both-fail.docx": {"status": "fail", "error": "BadZipFile"},
        "export/word-fails.docx": mine("anything"),
        "export/changed.docx": mine("New", sha256="bbb"),
    }

    export, imported = score(reading, sweep)

    totals = export.totals
    assert (totals["files"], totals["identical"]) == (2, 1)
    assert (totals["missing"], totals["extra"]) == (1, 1)  # "lost" missing; the label "1" extra
    assert (totals["parse_failures_word_ok"], totals["parse_failures_word_refuses"]) == (1, 1)
    assert export.skipped == Counter({"Word fail": 1, "changed since the Word reading": 1})
    assert imported.skipped == Counter({"not swept": 1})
    assert "export/fails.docx" in format_report([export, imported])
    assert report_json([export])["export"]["differing"][0]["missing"] == {"lost": 1}


def test_math_glyphs_are_counted_apart_from_other_missing_words():
    reading = {"meta": {}, "files": {"export/math.docx": word_record("Area 𝐴 = 𝜋 𝑟 2")}}

    export, _ = score(reading, {"export/math.docx": mine("Area", math=["A=\\pi r^{2}"])})

    assert export.totals["missing"] == 4
    assert export.totals["missing_math"] == 3


def test_paired_lists_status_changes_and_changed_words_only():
    reading = {
        "meta": {},
        "files": {
            "export/fixed.docx": word_record("Now read"),
            "export/better.docx": word_record("One two three"),
            "export/same.docx": word_record("Same"),
        },
    }
    before = {
        "export/fixed.docx": {"status": "fail", "error": "KeyError"},
        "export/better.docx": mine("One"),
        "export/same.docx": mine("Same"),
    }
    after = {
        "export/fixed.docx": mine("Now read"),
        "export/better.docx": mine("One two three 1."),
        "export/same.docx": mine("same"),
    }

    changes = {change.key: change for change in paired(reading, before, after)}

    assert set(changes) == {"export/fixed.docx", "export/better.docx"}
    assert (changes["export/fixed.docx"].before, changes["export/fixed.docx"].after) == ("fail", "ok")
    better = changes["export/better.docx"]
    assert (better.missing_delta, better.extra_delta) == (-2, 1)
    assert better.gained == Counter({"two": 1, "three": 1, "1": 1})
    assert "fail -> ok" in format_changes(list(changes.values()))
    assert format_changes([]) == "no file reads differently"


def test_a_checkout_is_required_and_checked(tmp_path, monkeypatch):
    monkeypatch.delenv(corpus_module.CORE_ENV, raising=False)
    with pytest.raises(CorpusError, match="--core"):
        core_root(None)
    with pytest.raises(CorpusError, match="not a LibreOffice core checkout"):
        core_root(str(tmp_path))


def fake_checkout(tmp_path):
    for directory in SETS.values():
        (tmp_path / directory).mkdir(parents=True)
    document = docx.Document()
    document.add_paragraph("Hello from the corpus.")
    buffer = io.BytesIO()
    document.save(buffer)
    (tmp_path / SETS["export"] / "hello.docx").write_bytes(buffer.getvalue())
    (tmp_path / SETS["import"] / "broken.docx").write_bytes(b"not a zip")
    (tmp_path / SETS["import"] / "notes.txt").write_text("ignored")
    return tmp_path


def test_a_sweep_reads_every_file_resumes_and_scores(tmp_path, monkeypatch):
    root = fake_checkout(tmp_path / "core")
    monkeypatch.setattr(sweep_module, "CACHE", tmp_path / "cache")
    assert [key for key, _ in corpus_files(root)] == ["export/hello.docx", "import/broken.docx"]

    out = sweep_module.sweep(root, "test", jobs=1)
    records = sweep_module.load_sweep(str(out))

    assert records["export/hello.docx"]["status"] == "ok"
    assert "Hello from the corpus." in records["export/hello.docx"]["text"]
    assert records["import/broken.docx"]["status"] == "fail"
    (out / cache_name("export/hello.docx")).write_text(json.dumps(mine("kept")), encoding="utf8")
    sweep_module.sweep(root, "test", jobs=1)
    assert sweep_module.load_sweep("test")["export/hello.docx"]["text"] == "kept"  # resumed, not redone


def test_the_word_reading_round_trips_byte_for_byte(tmp_path):
    reading = {"meta": {"date": "2026-01-01"}, "files": {"export/a.docx": word_record("Text é")}}
    first, second = tmp_path / "first.json.gz", tmp_path / "second.json.gz"

    save_word_reading(reading, first)
    save_word_reading(reading, second)

    assert first.read_bytes() == second.read_bytes()
    assert load_word_reading(first) == reading
    with pytest.raises(CorpusError, match="no Word reading"):
        load_word_reading(tmp_path / "absent.json.gz")


def test_the_committed_word_reading_is_complete():
    reading = load_word_reading()

    assert set(reading["meta"]) == {"date", "word_build", "libreoffice_commit"}
    assert len(reading["files"]) > 1500
    assert {key.split("/", 1)[0] for key in reading["files"]} == set(SETS)
    for key, record in reading["files"].items():
        assert record["status"] in {"ok", "fail", "hang"}, key
        assert len(record["sha256"]) == 64, key
        if record["status"] == "ok":
            assert isinstance(record["main"], str), key
