#  Copyright (c) 2025 Tom Villani, Ph.D.
#
# tests/integration/renderers/test_cross_renderer.py
"""Cross-renderer integration tests.

Tests cover:
- Cross-renderer consistency
- Complex scenarios across multiple renderers
- Error handling
- File output methods
- Real-world document structures
- Round-trip conversions

"""

import importlib.util

import pytest

DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None
REPORTLAB_AVAILABLE = importlib.util.find_spec("reportlab") is not None
EBOOKLIB_AVAILABLE = importlib.util.find_spec("ebooklib") is not None
PPTX_AVAILABLE = importlib.util.find_spec("pptx") is not None

from all2md.ast import (
    BlockQuote,
    CodeBlock,
    Document,
    Heading,
    Image,
    Link,
    List,
    ListItem,
    Paragraph,
    Strong,
    Table,
    TableCell,
    TableRow,
    Text,
    ThematicBreak,
)
from all2md.options import (
    HtmlRendererOptions,
)
from all2md.renderers import MarkdownRenderer
from all2md.renderers.html import HtmlRenderer

if DOCX_AVAILABLE:
    from all2md.renderers.docx import DocxRenderer

if REPORTLAB_AVAILABLE:
    from all2md.renderers.pdf import PdfRenderer

if EBOOKLIB_AVAILABLE:
    from all2md.renderers.epub import EpubRenderer

if PPTX_AVAILABLE:
    from pptx import Presentation

    from all2md.renderers.pptx import PptxRenderer


def create_sample_document():
    """Create a sample AST document for testing.

    Returns
    -------
    Document
        A sample document with various elements for testing.

    """
    return Document(
        metadata={"title": "Sample Document", "author": "Test Author"},
        children=[
            Heading(level=1, content=[Text(content="Document Title")]),
            Paragraph(
                content=[
                    Text(content="This is a paragraph with "),
                    Strong(content=[Text(content="bold text")]),
                    Text(content=" and a "),
                    Link(url="https://example.com", content=[Text(content="link")]),
                    Text(content="."),
                ]
            ),
            Heading(level=2, content=[Text(content="Lists")]),
            List(
                ordered=False,
                items=[
                    ListItem(children=[Paragraph(content=[Text(content="First item")])]),
                    ListItem(children=[Paragraph(content=[Text(content="Second item")])]),
                    ListItem(children=[Paragraph(content=[Text(content="Third item")])]),
                ],
            ),
            Heading(level=2, content=[Text(content="Code Example")]),
            CodeBlock(content='def hello():\n    print("Hello, world!")', language="python"),
            Heading(level=2, content=[Text(content="Table")]),
            Table(
                header=TableRow(
                    cells=[TableCell(content=[Text(content="Name")]), TableCell(content=[Text(content="Value")])]
                ),
                rows=[
                    TableRow(
                        cells=[TableCell(content=[Text(content="Alpha")]), TableCell(content=[Text(content="1")])]
                    ),
                    TableRow(cells=[TableCell(content=[Text(content="Beta")]), TableCell(content=[Text(content="2")])]),
                ],
            ),
            Heading(level=2, content=[Text(content="Quote")]),
            BlockQuote(children=[Paragraph(content=[Text(content="This is a blockquote.")])]),
        ],
    )


@pytest.mark.integration
class TestCrossRendererConsistency:
    """Tests for consistency across different renderers."""

    def test_same_document_all_renderers(self, tmp_path):
        """Test that same AST renders successfully to all formats."""
        doc = create_sample_document()

        # Markdown
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert len(md_result) > 0

        # HTML
        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert len(html_result) > 0

        # DOCX (if available)
        if DOCX_AVAILABLE:
            docx_renderer = DocxRenderer()
            docx_file = tmp_path / "consistency.docx"
            docx_renderer.render(doc, docx_file)
            assert docx_file.exists()

        # PDF (if available)
        if REPORTLAB_AVAILABLE:
            pdf_renderer = PdfRenderer()
            pdf_file = tmp_path / "consistency.pdf"
            pdf_renderer.render(doc, pdf_file)
            assert pdf_file.exists()

    def test_content_preservation(self):
        """Test that essential content is preserved across renderers."""
        doc = Document(
            children=[
                Heading(level=1, content=[Text(content="Test Title")]),
                Paragraph(content=[Text(content="Test content")]),
            ]
        )

        # Markdown
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "Test Title" in md_result
        assert "Test content" in md_result

        # HTML
        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert "Test Title" in html_result
        assert "Test content" in html_result


@pytest.mark.integration
class TestComplexScenarios:
    """Tests for complex real-world scenarios."""

    def test_deeply_nested_structures(self):
        """Test handling of deeply nested structures."""
        doc = Document(
            children=[
                List(
                    ordered=False,
                    items=[
                        ListItem(
                            children=[
                                Paragraph(content=[Text(content="Item 1")]),
                                List(
                                    ordered=False,
                                    items=[ListItem(children=[Paragraph(content=[Text(content="Nested 1")])])],
                                ),
                            ]
                        )
                    ],
                )
            ]
        )

        # Should render without errors
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "Item 1" in md_result
        assert "Nested 1" in md_result

        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert "Item 1" in html_result
        assert "Nested 1" in html_result

    def test_mixed_formatting(self):
        """Test paragraph with multiple formatting types."""
        doc = Document(
            children=[
                Paragraph(
                    content=[
                        Text(content="This is "),
                        Strong(content=[Text(content="bold")]),
                        Text(content=" and this is "),
                        Link(url="http://example.com", content=[Text(content="a link")]),
                        Text(content="."),
                    ]
                )
            ]
        )

        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "**bold**" in md_result
        assert "[a link](http://example.com)" in md_result

        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert "<strong>bold</strong>" in html_result
        assert '<a href="http://example.com">a link</a>' in html_result

    def test_large_table(self, tmp_path):
        """Test handling of large tables."""
        # Create table with many rows
        rows = [
            TableRow(cells=[TableCell(content=[Text(content=f"Cell {i}-{j}")]) for j in range(5)]) for i in range(20)
        ]

        doc = Document(
            children=[
                Table(
                    header=TableRow(cells=[TableCell(content=[Text(content=f"Col {i}")]) for i in range(5)]), rows=rows
                )
            ]
        )

        # Should handle large tables
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "Cell 0-0" in md_result
        assert "Cell 19-4" in md_result

        if DOCX_AVAILABLE:
            docx_renderer = DocxRenderer()
            docx_file = tmp_path / "large_table.docx"
            docx_renderer.render(doc, docx_file)
            assert docx_file.exists()


@pytest.mark.integration
class TestErrorHandling:
    """Tests for error handling and edge cases."""

    def test_empty_document_all_renderers(self, tmp_path):
        """Test that empty documents are handled gracefully."""
        doc = Document()

        # Markdown
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert md_result == ""

        # HTML
        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert html_result == ""

        # DOCX (if available)
        if DOCX_AVAILABLE:
            docx_renderer = DocxRenderer()
            docx_file = tmp_path / "empty.docx"
            docx_renderer.render(doc, docx_file)
            assert docx_file.exists()

        # PDF (if available)
        if REPORTLAB_AVAILABLE:
            pdf_renderer = PdfRenderer()
            pdf_file = tmp_path / "empty.pdf"
            pdf_renderer.render(doc, pdf_file)
            assert pdf_file.exists()

    def test_missing_optional_attributes(self):
        """Test handling of missing optional attributes."""
        # Image without title
        doc = Document(children=[Paragraph(content=[Image(url="image.png", alt_text="Image")])])

        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "![Image](image.png)" in md_result

        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert '<img src="image.png" alt="Image">' in html_result


@pytest.mark.integration
class TestFileOutput:
    """Tests for file output methods."""

    def test_markdown_to_file(self, tmp_path):
        """Test Markdown rendering to file."""
        doc = create_sample_document()
        renderer = MarkdownRenderer()
        output_file = tmp_path / "output.md"
        renderer.render(doc, output_file)

        assert output_file.exists()
        content = output_file.read_text()
        assert "Document Title" in content

    def test_html_to_file(self, tmp_path):
        """Test HTML rendering to file."""
        doc = create_sample_document()
        renderer = HtmlRenderer(HtmlRendererOptions(standalone=True))
        output_file = tmp_path / "output.html"
        renderer.render(doc, output_file)

        assert output_file.exists()
        content = output_file.read_text()
        assert "<!DOCTYPE html>" in content


@pytest.mark.integration
class TestRealWorldDocuments:
    """Tests with realistic document structures."""

    def test_technical_documentation(self, tmp_path):
        """Test rendering technical documentation structure."""
        doc = Document(
            children=[
                Heading(level=1, content=[Text(content="API Documentation")]),
                Heading(level=2, content=[Text(content="Overview")]),
                Paragraph(content=[Text(content="This API provides access to our services.")]),
                Heading(level=2, content=[Text(content="Endpoints")]),
                Heading(level=3, content=[Text(content="GET /users")]),
                Paragraph(content=[Text(content="Retrieves all users.")]),
                CodeBlock(content="GET /api/users\nAuthorization: Bearer token", language="http"),
                Heading(level=3, content=[Text(content="Response")]),
                CodeBlock(content='{\n  "users": []\n}', language="json"),
            ]
        )

        # Should render successfully to all formats
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "# API Documentation" in md_result

        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=True, include_toc=True))
        html_result = html_renderer.render_to_string(doc)
        assert "API Documentation" in html_result

        if DOCX_AVAILABLE:
            docx_renderer = DocxRenderer()
            docx_file = tmp_path / "api_docs.docx"
            docx_renderer.render(doc, docx_file)
            assert docx_file.exists()

    def test_blog_post_structure(self):
        """Test rendering blog post structure."""
        doc = Document(
            metadata={"title": "My Blog Post", "author": "John Doe"},
            children=[
                Heading(level=1, content=[Text(content="My Blog Post")]),
                Paragraph(content=[Text(content="Posted on January 1, 2025")]),
                Paragraph(content=[Text(content="This is the introduction to my blog post.")]),
                Heading(level=2, content=[Text(content="Main Point")]),
                Paragraph(content=[Text(content="Here's the main content.")]),
                BlockQuote(children=[Paragraph(content=[Text(content="A relevant quote.")])]),
                Heading(level=2, content=[Text(content="Conclusion")]),
                Paragraph(content=[Text(content="Final thoughts.")]),
            ],
        )

        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert "My Blog Post" in md_result

        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=True))
        html_result = html_renderer.render_to_string(doc)
        assert "My Blog Post" in html_result
        assert "<blockquote>" in html_result


@pytest.mark.skipif(not (EBOOKLIB_AVAILABLE and PPTX_AVAILABLE), reason="ebooklib or python-pptx not installed")
@pytest.mark.integration
class TestNewRendererConsistency:
    """Test consistency between new renderers and existing ones."""

    def test_same_ast_all_renderers(self, tmp_path):
        """Test that same AST renders successfully to all formats including new ones."""
        doc = create_sample_document()

        # Markdown
        md_renderer = MarkdownRenderer()
        md_result = md_renderer.render_to_string(doc)
        assert len(md_result) > 0

        # HTML
        html_renderer = HtmlRenderer(HtmlRendererOptions(standalone=False))
        html_result = html_renderer.render_to_string(doc)
        assert len(html_result) > 0

        # DOCX (if available)
        if DOCX_AVAILABLE:
            docx_renderer = DocxRenderer()
            docx_file = tmp_path / "all_formats.docx"
            docx_renderer.render(doc, docx_file)
            assert docx_file.exists()

        # PDF (if available)
        if REPORTLAB_AVAILABLE:
            pdf_renderer = PdfRenderer()
            pdf_file = tmp_path / "all_formats.pdf"
            pdf_renderer.render(doc, pdf_file)
            assert pdf_file.exists()

        # EPUB
        epub_renderer = EpubRenderer()
        epub_file = tmp_path / "all_formats.epub"
        epub_renderer.render(doc, epub_file)
        assert epub_file.exists()

        # PPTX
        pptx_renderer = PptxRenderer()
        pptx_file = tmp_path / "all_formats.pptx"
        pptx_renderer.render(doc, pptx_file)
        assert pptx_file.exists()

    def test_markdown_to_epub_to_markdown_roundtrip(self, tmp_path):
        """Test parsing EPUB and rendering back to Markdown."""
        from all2md.parsers.epub import EpubToAstConverter

        # Create EPUB from AST
        doc = Document(
            children=[
                Heading(level=1, content=[Text(content="Chapter 1")]),
                Paragraph(content=[Text(content="Some content here.")]),
            ]
        )

        epub_renderer = EpubRenderer()
        epub_file = tmp_path / "test.epub"
        epub_renderer.render(doc, epub_file)

        # Parse EPUB back to AST
        epub_parser = EpubToAstConverter()
        parsed_doc = epub_parser.parse(epub_file)

        # Render to Markdown
        md_renderer = MarkdownRenderer()
        result = md_renderer.render_to_string(parsed_doc)

        # Verify content preserved
        assert "Chapter 1" in result
        assert "Some content" in result

    def test_markdown_to_pptx_conversion(self, tmp_path):
        """Test converting Markdown-style document to PPTX."""
        # Create document with slide-like structure
        doc = Document(
            children=[
                Heading(level=1, content=[Text(content="Presentation")]),
                Paragraph(content=[Text(content="Welcome")]),
                ThematicBreak(),
                Heading(level=2, content=[Text(content="Agenda")]),
                List(
                    ordered=True,
                    items=[
                        ListItem(children=[Paragraph(content=[Text(content="Introduction")])]),
                        ListItem(children=[Paragraph(content=[Text(content="Main Points")])]),
                        ListItem(children=[Paragraph(content=[Text(content="Conclusion")])]),
                    ],
                ),
            ]
        )

        # Render to PPTX
        renderer = PptxRenderer()
        output_file = tmp_path / "converted.pptx"
        renderer.render(doc, output_file)

        # Verify structure
        prs = Presentation(str(output_file))
        assert len(prs.slides) == 2  # Two slides separated by ThematicBreak
